import math
import os
import tkinter as tk
from datetime import datetime
from tkinter import messagebox, ttk

from Calculadora import (
    MARGEN_OBJETIVO_PORCENTAJE,
    calcular_indicadores,
    clasificar_eficiencia,
    generar_recomendaciones,
    guardar_historial_xlsx,
    limitar_porcentaje,
    obtener_ruta_historial,
)


INDICADORES = [
    "Productividad (%)",
    "Eficiencia de costos (%)",
    "Calidad (%)",
    "Cumplimiento de entregas (%)",
    "Utilizacion de recursos (%)",
    "Margen operativo (%)",
]

ETIQUETAS_GRAFICO = [
    "Productividad",
    "Costos",
    "Calidad",
    "Entregas",
    "Recursos",
    "Margen",
]

GRUPOS_CAMPOS = [
    (
        "Producción",
        [
            ("Producción real (unidades)", "Produccion real"),
            ("Producción objetivo (unidades)", "Produccion objetivo"),
        ],
    ),
    (
        "Calidad",
        [
            ("Unidades producidas", "Unidades producidas"),
            ("Unidades conformes", "Unidades conformes"),
        ],
    ),
    (
        "Utilización de recursos",
        [
            ("Horas disponibles", "Horas disponibles"),
            ("Horas productivas", "Horas productivas"),
        ],
    ),
    (
        "Costos",
        [
            ("Costo presupuestado (COP)", "Costo presupuestado"),
            ("Costo real (COP)", "Costo real"),
        ],
    ),
    (
        "Entregas",
        [
            ("Total de entregas", "Entregas totales"),
            ("Entregas a tiempo", "Entregas a tiempo"),
        ],
    ),
    (
        "Finanzas",
        [
            ("Ventas totales (COP)", "Ventas totales"),
            ("Costos operativos (COP)", "Costos operativos"),
        ],
    ),
]


class CalculadoraDesktop:
    def __init__(self, root):
        self.root = root
        self.root.title("Calculadora de eficiencia empresarial")
        self.root.geometry("1100x780")
        self.root.minsize(920, 680)
        self.root.report_callback_exception = self.report_callback_exception

        self.campos = {}
        self.registros = []
        self.registro_comparado = None
        self.nuevo_historial_var = tk.BooleanVar(value=False)

        self.crear_interfaz()
        self.actualizar_historial()

    def crear_interfaz(self):
        contenedor = ttk.Frame(self.root, padding=14)
        contenedor.pack(fill="both", expand=True)
        contenedor.columnconfigure(0, weight=1)
        contenedor.rowconfigure(1, weight=1)

        ttk.Label(
            contenedor,
            text="Calculadora de eficiencia empresarial",
            font=("Segoe UI", 18, "bold"),
        ).grid(row=0, column=0, sticky="w", pady=(0, 10))

        self.pestanas = ttk.Notebook(contenedor)
        self.pestanas.grid(row=1, column=0, sticky="nsew")

        self.pestana_nueva = ttk.Frame(self.pestanas, padding=12)
        self.pestana_resultados = ttk.Frame(self.pestanas, padding=12)
        self.pestana_historial = ttk.Frame(self.pestanas, padding=12)
        self.pestana_comparar = ttk.Frame(self.pestanas, padding=12)

        self.pestanas.add(self.pestana_nueva, text="Nueva evaluación")
        self.pestanas.add(self.pestana_resultados, text="Resultados")
        self.pestanas.add(self.pestana_historial, text="Historial")
        self.pestanas.add(self.pestana_comparar, text="Comparar")

        self.crear_formulario()
        self.crear_pestana_resultados()
        self.crear_pestana_historial()
        self.crear_pestana_comparar()

        self.estado_var = tk.StringVar(
            value=f"Historial: {obtener_ruta_historial()}"
        )
        ttk.Label(
            contenedor,
            textvariable=self.estado_var,
            anchor="w",
            wraplength=1040,
        ).grid(row=2, column=0, sticky="ew", pady=(8, 0))

    def crear_formulario(self):
        identidad = ttk.LabelFrame(
            self.pestana_nueva, text="Identificación", padding=10
        )
        identidad.pack(fill="x", pady=(0, 10))
        identidad.columnconfigure(1, weight=1)
        identidad.columnconfigure(3, weight=1)

        ttk.Label(identidad, text="Empresa").grid(
            row=0, column=0, sticky="w", padx=(0, 6), pady=4
        )
        self.campos["Empresa"] = tk.StringVar()
        ttk.Entry(identidad, textvariable=self.campos["Empresa"]).grid(
            row=0, column=1, sticky="ew", padx=(0, 16), pady=4
        )

        ttk.Label(identidad, text="Periodo (AAAA-MM-DD)").grid(
            row=0, column=2, sticky="w", padx=(0, 6), pady=4
        )
        self.campos["Periodo"] = tk.StringVar(
            value=datetime.now().strftime("%Y-%m-%d")
        )
        ttk.Entry(identidad, textvariable=self.campos["Periodo"]).grid(
            row=0, column=3, sticky="ew", pady=4
        )

        grilla = ttk.Frame(self.pestana_nueva)
        grilla.pack(fill="both", expand=True)
        grilla.columnconfigure(0, weight=1)
        grilla.columnconfigure(1, weight=1)

        for indice, (titulo, campos) in enumerate(GRUPOS_CAMPOS):
            columna = indice // 3
            fila = indice % 3
            grupo = ttk.LabelFrame(grilla, text=titulo, padding=8)
            grupo.grid(
                row=fila,
                column=columna,
                sticky="nsew",
                padx=(0, 8) if columna == 0 else (8, 0),
                pady=5,
            )
            grupo.columnconfigure(1, weight=1)

            for numero, (etiqueta, clave) in enumerate(campos):
                variable = tk.StringVar()
                self.campos[clave] = variable
                ttk.Label(grupo, text=etiqueta).grid(
                    row=numero, column=0, sticky="w", padx=(0, 8), pady=4
                )
                ttk.Entry(grupo, textvariable=variable, width=20).grid(
                    row=numero, column=1, sticky="ew", pady=4
                )

        acciones = ttk.Frame(self.pestana_nueva)
        acciones.pack(fill="x", pady=(10, 0))
        ttk.Checkbutton(
            acciones,
            text="Iniciar un historial nuevo y archivar el actual",
            variable=self.nuevo_historial_var,
        ).pack(side="left")
        ttk.Button(
            acciones,
            text="Limpiar campos",
            command=self.limpiar_campos,
        ).pack(side="right", padx=(8, 0))
        ttk.Button(
            acciones,
            text="Calcular y guardar",
            command=self.calcular_y_guardar,
        ).pack(side="right")

        grilla.rowconfigure(0, weight=1)
        grilla.rowconfigure(1, weight=1)
        grilla.rowconfigure(2, weight=1)

    def crear_pestana_resultados(self):
        self.clasificacion_var = tk.StringVar(
            value="Los resultados aparecerán después del cálculo."
        )
        ttk.Label(
            self.pestana_resultados,
            textvariable=self.clasificacion_var,
            font=("Segoe UI", 15, "bold"),
        ).pack(anchor="w", pady=(0, 10))

        self.resultados_tabla = ttk.Treeview(
            self.pestana_resultados,
            columns=("indicador", "valor"),
            show="headings",
            height=8,
        )
        self.resultados_tabla.heading("indicador", text="Indicador")
        self.resultados_tabla.heading("valor", text="Resultado")
        self.resultados_tabla.column("indicador", width=390, anchor="w")
        self.resultados_tabla.column("valor", width=160, anchor="e")
        self.resultados_tabla.pack(fill="x", pady=(0, 12))

        ttk.Label(
            self.pestana_resultados,
            text="Recomendaciones",
            font=("Segoe UI", 11, "bold"),
        ).pack(anchor="w")
        self.recomendaciones_texto = tk.Text(
            self.pestana_resultados, height=8, wrap="word", state="disabled"
        )
        self.recomendaciones_texto.pack(fill="both", expand=True, pady=(5, 0))

    def crear_pestana_historial(self):
        barra = ttk.Frame(self.pestana_historial)
        barra.pack(fill="x", pady=(0, 8))
        ttk.Label(
            barra,
            text="Evaluaciones guardadas en el archivo de historial activo.",
        ).pack(side="left")
        ttk.Button(
            barra, text="Actualizar", command=self.actualizar_historial
        ).pack(side="right")

        columnas = ("empresa", "periodo", "indice", "clasificacion")
        self.historial_tabla = ttk.Treeview(
            self.pestana_historial,
            columns=columnas,
            show="headings",
            selectmode="browse",
        )
        for clave, titulo, ancho, ancla in [
            ("empresa", "Empresa", 240, "w"),
            ("periodo", "Periodo", 140, "center"),
            ("indice", "Índice global", 140, "e"),
            ("clasificacion", "Clasificación", 230, "w"),
        ]:
            self.historial_tabla.heading(clave, text=titulo)
            self.historial_tabla.column(clave, width=ancho, anchor=ancla)
        self.historial_tabla.pack(fill="both", expand=True)

    def crear_pestana_comparar(self):
        ttk.Label(
            self.pestana_comparar,
            text=(
                "Selecciona exactamente dos evaluaciones del historial "
                "y pulsa «Comparar seleccionadas»."
            ),
            wraplength=1000,
        ).pack(anchor="w", pady=(0, 8))

        columnas = ("empresa", "periodo", "indice", "clasificacion")
        self.comparar_tabla = ttk.Treeview(
            self.pestana_comparar,
            columns=columnas,
            show="headings",
            selectmode="extended",
            height=7,
        )
        for clave, titulo, ancho, ancla in [
            ("empresa", "Empresa", 240, "w"),
            ("periodo", "Periodo", 140, "center"),
            ("indice", "Índice global", 140, "e"),
            ("clasificacion", "Clasificación", 230, "w"),
        ]:
            self.comparar_tabla.heading(clave, text=titulo)
            self.comparar_tabla.column(clave, width=ancho, anchor=ancla)
        self.comparar_tabla.pack(fill="x", expand=False)

        ttk.Button(
            self.pestana_comparar,
            text="Comparar seleccionadas",
            command=self.comparar_seleccionadas,
        ).pack(anchor="e", pady=8)

        self.comparacion_tabla = ttk.Treeview(
            self.pestana_comparar,
            columns=("indicador", "primera", "segunda"),
            show="headings",
            height=7,
        )
        self.comparacion_tabla.heading("indicador", text="Indicador")
        self.comparacion_tabla.heading("primera", text="Evaluación 1")
        self.comparacion_tabla.heading("segunda", text="Evaluación 2")
        self.comparacion_tabla.column("indicador", width=330, anchor="w")
        self.comparacion_tabla.column("primera", width=180, anchor="e")
        self.comparacion_tabla.column("segunda", width=180, anchor="e")
        self.comparacion_tabla.pack(fill="x", pady=(0, 8))

        self.grafico = tk.Canvas(
            self.pestana_comparar,
            height=250,
            background="white",
            highlightthickness=1,
            highlightbackground="#cccccc",
        )
        self.grafico.pack(fill="both", expand=True)
        self.grafico.bind("<Configure>", self.redibujar_grafico)

    def leer_decimal(self, clave, etiqueta, positivo=False):
        texto = self.campos[clave].get().strip().replace(",", ".")
        if not texto:
            raise ValueError(f"El campo «{etiqueta}» es obligatorio.")
        try:
            valor = float(texto)
        except ValueError as error:
            raise ValueError(
                f"Escribe un número válido en «{etiqueta}»."
            ) from error
        if not math.isfinite(valor):
            raise ValueError(f"«{etiqueta}» debe ser un número finito.")
        if valor < 0 or (positivo and valor <= 0):
            condicion = "mayor que cero" if positivo else "cero o mayor"
            raise ValueError(f"«{etiqueta}» debe ser {condicion}.")
        return valor

    def leer_entero(self, clave, etiqueta, minimo=0):
        texto = self.campos[clave].get().strip()
        if not texto:
            raise ValueError(f"El campo «{etiqueta}» es obligatorio.")
        try:
            valor = int(texto)
        except ValueError as error:
            raise ValueError(
                f"Escribe un número entero válido en «{etiqueta}»."
            ) from error
        if valor < minimo:
            raise ValueError(
                f"«{etiqueta}» debe ser como mínimo {minimo}."
            )
        return valor

    def recopilar_datos(self):
        empresa = self.campos["Empresa"].get().strip()
        if not empresa:
            raise ValueError("El nombre de la empresa es obligatorio.")

        periodo = self.campos["Periodo"].get().strip()
        if not periodo:
            periodo = datetime.now().strftime("%Y-%m-%d")
        try:
            datetime.strptime(periodo, "%Y-%m-%d")
        except ValueError as error:
            raise ValueError(
                "El periodo debe tener el formato AAAA-MM-DD, "
                "por ejemplo, 2026-10-01."
            ) from error

        datos = {
            "Empresa": empresa,
            "Periodo": periodo,
            "Produccion real": self.leer_decimal(
                "Produccion real", "Producción real"
            ),
            "Produccion objetivo": self.leer_decimal(
                "Produccion objetivo", "Producción objetivo", positivo=True
            ),
            "Costo presupuestado": self.leer_decimal(
                "Costo presupuestado", "Costo presupuestado"
            ),
            "Costo real": self.leer_decimal("Costo real", "Costo real"),
            "Unidades producidas": self.leer_entero(
                "Unidades producidas", "Unidades producidas", minimo=1
            ),
            "Unidades conformes": self.leer_entero(
                "Unidades conformes", "Unidades conformes"
            ),
            "Entregas totales": self.leer_entero(
                "Entregas totales", "Total de entregas", minimo=1
            ),
            "Entregas a tiempo": self.leer_entero(
                "Entregas a tiempo", "Entregas a tiempo"
            ),
            "Horas disponibles": self.leer_decimal(
                "Horas disponibles", "Horas disponibles", positivo=True
            ),
            "Horas productivas": self.leer_decimal(
                "Horas productivas", "Horas productivas"
            ),
            "Ventas totales": self.leer_decimal(
                "Ventas totales", "Ventas totales", positivo=True
            ),
            "Costos operativos": self.leer_decimal(
                "Costos operativos", "Costos operativos"
            ),
        }

        if datos["Unidades conformes"] > datos["Unidades producidas"]:
            raise ValueError(
                "Las unidades conformes no pueden superar las producidas."
            )
        if datos["Entregas a tiempo"] > datos["Entregas totales"]:
            raise ValueError(
                "Las entregas a tiempo no pueden superar el total."
            )
        if datos["Horas productivas"] > datos["Horas disponibles"]:
            raise ValueError(
                "Las horas productivas no pueden superar las disponibles."
            )

        return datos

    def calcular_y_guardar(self):
        try:
            datos = self.recopilar_datos()
            indicadores = calcular_indicadores(datos)
            clasificacion = clasificar_eficiencia(
                indicadores["Indice global (%)"]
            )
            recomendaciones = generar_recomendaciones(indicadores)

            registro = datos.copy()
            registro.update(indicadores)
            registro["Clasificacion"] = clasificacion
            registro["Recomendaciones"] = " | ".join(recomendaciones)

            ruta, ruta_archivado = guardar_historial_xlsx(
                registro,
                iniciar_nuevo=self.nuevo_historial_var.get(),
            )
        except ValueError as error:
            messagebox.showerror("Revisa los datos", str(error), parent=self.root)
            return
        except PermissionError:
            messagebox.showerror(
                "No se pudo guardar",
                "Cierra el archivo Historial_Eficiencia.xlsx en Excel "
                "y vuelve a intentarlo. No se reemplazó el historial.",
                parent=self.root,
            )
            return
        except OSError as error:
            messagebox.showerror(
                "No se pudo guardar",
                f"No se pudo guardar el historial.\n\n{error}",
                parent=self.root,
            )
            return

        self.nuevo_historial_var.set(False)
        self.mostrar_resultados(indicadores, clasificacion, recomendaciones)
        self.actualizar_historial()

        if ruta_archivado:
            self.estado_var.set(
                f"Historial anterior archivado en: {ruta_archivado}  |  "
                f"Nuevo historial: {ruta}"
            )
        else:
            self.estado_var.set(f"Evaluación guardada en: {ruta}")
        self.pestanas.select(self.pestana_resultados)

    def mostrar_resultados(self, indicadores, clasificacion, recomendaciones):
        self.clasificacion_var.set(
            f"{clasificacion} — Índice global: "
            f"{indicadores['Indice global (%)']:.2f}%"
        )
        for item in self.resultados_tabla.get_children():
            self.resultados_tabla.delete(item)

        for nombre, valor in indicadores.items():
            self.resultados_tabla.insert(
                "", "end", values=(nombre, f"{valor:.2f}%")
            )

        self.recomendaciones_texto.configure(state="normal")
        self.recomendaciones_texto.delete("1.0", tk.END)
        self.recomendaciones_texto.insert(
            "1.0",
            "\n".join(
                f"{numero}. {texto}"
                for numero, texto in enumerate(recomendaciones, start=1)
            ),
        )
        self.recomendaciones_texto.configure(state="disabled")

    def actualizar_historial(self):
        self.registros = []
        ruta = obtener_ruta_historial()
        if os.path.exists(ruta):
            from openpyxl import load_workbook

            libro = load_workbook(ruta, read_only=True, data_only=True)
            try:
                if "Historial" in libro.sheetnames:
                    hoja = libro["Historial"]
                    filas = hoja.iter_rows(values_only=True)
                    encabezados = next(filas, ())
                    for fila in filas:
                        if not fila or all(valor is None for valor in fila):
                            continue
                        registro = {
                            encabezado: fila[indice]
                            for indice, encabezado in enumerate(encabezados)
                            if encabezado is not None and indice < len(fila)
                        }
                        self.registros.append(registro)
            finally:
                libro.close()

        self.llenar_tablas_historial()

    def llenar_tablas_historial(self):
        for tabla in (self.historial_tabla, self.comparar_tabla):
            for item in tabla.get_children():
                tabla.delete(item)

        for indice, registro in enumerate(self.registros):
            valores = (
                registro.get("Empresa", ""),
                registro.get("Periodo", ""),
                self.formatear_porcentaje(registro.get("Indice global (%)")),
                registro.get("Clasificacion", ""),
            )
            self.historial_tabla.insert("", "end", values=valores)
            self.comparar_tabla.insert(
                "", "end", iid=str(indice), values=valores
            )

    @staticmethod
    def formatear_porcentaje(valor):
        try:
            return f"{float(valor):.2f}%"
        except (TypeError, ValueError):
            return ""

    def comparar_seleccionadas(self):
        seleccion = self.comparar_tabla.selection()
        if len(seleccion) != 2:
            messagebox.showinfo(
                "Selecciona dos evaluaciones",
                "Para comparar, selecciona exactamente dos filas.",
                parent=self.root,
            )
            return

        primera, segunda = (
            self.registros[int(indice)] for indice in seleccion
        )
        self.registro_comparado = (primera, segunda)

        for item in self.comparacion_tabla.get_children():
            self.comparacion_tabla.delete(item)
        for nombre in INDICADORES + ["Indice global (%)"]:
            self.comparacion_tabla.insert(
                "",
                "end",
                values=(
                    nombre,
                    self.formatear_porcentaje(primera.get(nombre)),
                    self.formatear_porcentaje(segunda.get(nombre)),
                ),
            )

        self.redibujar_grafico()

    @staticmethod
    def puntajes_normalizados(registro):
        puntajes = {
            clave: limitar_porcentaje(
                float(registro.get(clave, 0) or 0)
            )
            for clave in INDICADORES[:-1]
        }
        margen = float(registro.get("Margen operativo (%)", 0) or 0)
        puntajes["Margen operativo (%)"] = limitar_porcentaje(
            margen / MARGEN_OBJETIVO_PORCENTAJE * 100
        )
        return [puntajes[nombre] for nombre in INDICADORES]

    def redibujar_grafico(self, _evento=None):
        canvas = self.grafico
        canvas.delete("all")
        if not self.registro_comparado:
            canvas.create_text(
                12,
                18,
                anchor="nw",
                text="Selecciona dos evaluaciones para ver el gráfico.",
                fill="#555555",
            )
            return

        primera, segunda = self.registro_comparado
        puntajes_a = self.puntajes_normalizados(primera)
        puntajes_b = self.puntajes_normalizados(segunda)
        ancho = max(canvas.winfo_width(), 500)
        alto = max(canvas.winfo_height(), 210)
        margen_izquierdo = 46
        margen_derecho = 18
        margen_superior = 42
        margen_inferior = 48
        alto_grafico = alto - margen_superior - margen_inferior
        ancho_grafico = ancho - margen_izquierdo - margen_derecho
        base_y = margen_superior + alto_grafico

        for valor in (0, 50, 100):
            y = base_y - alto_grafico * valor / 100
            canvas.create_line(
                margen_izquierdo, y, ancho - margen_derecho, y,
                fill="#dddddd",
            )
            canvas.create_text(
                margen_izquierdo - 8, y, text=str(valor), anchor="e",
                fill="#555555",
            )

        canvas.create_text(
            margen_izquierdo,
            14,
            anchor="nw",
            text="Puntajes normalizados (0–100, según el índice global)",
            fill="#333333",
        )

        ancho_grupo = ancho_grafico / len(ETIQUETAS_GRAFICO)
        ancho_barra = min(24, ancho_grupo * 0.24)
        colores = ("#2878b5", "#e67e22")
        etiqueta_a = (
            f"{primera.get('Empresa', '')} "
            f"({primera.get('Periodo', '')})"
        )
        etiqueta_b = (
            f"{segunda.get('Empresa', '')} "
            f"({segunda.get('Periodo', '')})"
        )
        canvas.create_text(
            margen_izquierdo + 20, 30, anchor="w",
            text=etiqueta_a, fill=colores[0],
        )
        canvas.create_text(
            margen_izquierdo + ancho_grafico * 0.52, 30, anchor="w",
            text=etiqueta_b, fill=colores[1],
        )

        for indice, etiqueta in enumerate(ETIQUETAS_GRAFICO):
            centro = margen_izquierdo + ancho_grupo * (indice + 0.5)
            for desplazamiento, valor, color in [
                (-ancho_barra / 2, puntajes_a[indice], colores[0]),
                (ancho_barra / 2, puntajes_b[indice], colores[1]),
            ]:
                x1 = centro + desplazamiento
                x2 = x1 + ancho_barra * 0.9
                y1 = base_y - alto_grafico * valor / 100
                canvas.create_rectangle(
                    x1, y1, x2, base_y, fill=color, outline=""
                )
            canvas.create_text(
                centro,
                base_y + 20,
                text=etiqueta,
                width=max(60, ancho_grupo - 4),
                justify="center",
                fill="#333333",
            )

    def limpiar_campos(self):
        for clave, variable in self.campos.items():
            if clave in ("Empresa", "Periodo"):
                continue
            variable.set("")
        self.campos["Empresa"].set("")
        self.campos["Periodo"].set(datetime.now().strftime("%Y-%m-%d"))

    def report_callback_exception(self, exc_type, exc_value, _traceback):
        messagebox.showerror(
            "Se produjo un error",
            f"{exc_type.__name__}: {exc_value}",
            parent=self.root,
        )


def main():
    root = tk.Tk()
    CalculadoraDesktop(root)
    root.mainloop()


if __name__ == "__main__":
    main()
